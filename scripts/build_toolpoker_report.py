from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "research" / "ToolPoker-PokerCoach-Product-Recommendations.docx"

BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(31, 77, 120)
INK = RGBColor(11, 37, 69)
GRAY = RGBColor(90, 99, 110)
LIGHT = "F2F4F7"
PALE_BLUE = "E8EEF5"


def set_font(run, size=11, color=None, bold=None, italic=None, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_doc(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK, 8, 4),
    ):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    for lname in ("List Bullet", "List Number", "List Number 2"):
        st = doc.styles[lname]
        st.font.name = "Calibri"
        st.font.size = Pt(11)
        st.paragraph_format.left_indent = Inches(0.5)
        st.paragraph_format.first_line_indent = Inches(-0.25)
        st.paragraph_format.space_after = Pt(8)
        st.paragraph_format.line_spacing = 1.167


def add_header_footer(doc):
    section = doc.sections[0]
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(2)
    set_font(hp.add_run("POKERCOACH • RESEARCH BRIEF"), 9, GRAY, True)
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(fp.add_run("ToolPoker → PokerCoach  |  25 August 2026"), 9, GRAY)


def add_title(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run("PRODUCT & ARCHITECTURE BRIEF"), 10, BLUE, True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    set_font(p.add_run("ToolPoker → PokerCoach"), 24, INK, True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    set_font(p.add_run("How agentic solver grounding can improve poker hand analysis"), 14, GRAY)
    for label, value in (
        ("Audience", "PokerCoach product and engineering"),
        ("Decision", "Prioritize unified evidence + semantic decision contracts; defer fine-tuning"),
        ("Evidence", "arXiv:2602.00528, CFR/DeepCFR, PokerBench, and current repository inspection"),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        set_font(p.add_run(f"{label}: "), 10.5, INK, True)
        set_font(p.add_run(value), 10.5, INK)


def add_callout(doc, label, text, fill=PALE_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run(label + "  "), 11, INK, True)
    set_font(p.add_run(text), 11, INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_p(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        set_font(p.add_run(bold_lead), 11, INK, True)
        set_font(p.add_run(text[len(bold_lead):]), 11, INK)
    else:
        set_font(p.add_run(text), 11, INK)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_font(p.add_run(item), 11, INK)


def add_numbered(doc, items, style="List Number"):
    for item in items:
        p = doc.add_paragraph(style=style)
        set_font(p.add_run(item), 11, INK)


def add_link(p, label, url):
    hyperlink = OxmlElement("w:hyperlink")
    rel_id = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "2E74B5")
    underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single")
    rpr.extend([color, underline]); run.append(rpr)
    text = OxmlElement("w:t"); text.text = label; run.append(text)
    hyperlink.append(run); p._p.append(hyperlink)


def heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def build():
    doc = Document()
    style_doc(doc)
    add_header_footer(doc)
    add_title(doc)
    add_callout(doc, "Executive answer", "Yes. Make one validated evidence packet the source of truth for both the poker recommendation and its explanation. PokerCoach already has most building blocks, but they are split across parallel paths. Do not begin with model fine-tuning.")

    heading(doc, "The paper’s central lesson", 1)
    add_p(doc, "Lin et al. identify three recurring LLM poker failures: heuristic shortcuts, factual misunderstandings, and a knowing–doing gap where the final action contradicts the reasoning. Behavior cloning plus regret-inspired reinforcement learning improved professional-sounding reasoning, but did not produce accurate equilibrium play.")
    add_p(doc, "ToolPoker changes the architecture: a unified tool call returns the solver action together with equity, pot odds, and range information, then the LLM explains those outputs. In the paper’s evaluation, this closes far more of the gameplay and factual-alignment gap than reasoning-style training alone.")
    p = doc.add_paragraph(); set_font(p.add_run("Primary paper: "), 10, GRAY, True); add_link(p, "Lin et al., arXiv:2602.00528", "https://arxiv.org/html/2602.00528")

    heading(doc, "What maps directly to PokerCoach", 1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ["Paper method / finding", "PokerCoach today", "Recommended move"]
    for i, text in enumerate(headers):
        shade(table.rows[0].cells[i], LIGHT)
        p = table.rows[0].cells[i].paragraphs[0]; set_font(p.add_run(text), 9.5, INK, True)
    rows = [
        ("Unified solver interface", "Local math, LLM panels, range interpretation, and TexasSolver run as separate paths.", "Return one immutable evidence packet used by every panel."),
        ("State mis-specification is a residual ToolPoker error", "Screenshot import and repaired actions/cards create more state risk than a clean simulator.", "Add server-side solver preflight, field provenance, and a canonical state hash."),
        ("Action alignment dominates the reward ablation", "Solver output is an optional sibling check; the primary recommendation may already be displayed.", "Use a deterministic decision contract; LLM explains but cannot override a valid solver mix."),
        ("Factual alignment improves with tools", "Strategic modes omit equity/range/EV because local values may be unstable.", "Promote verified math and solver quantities; reject uncited numbers."),
        ("Compact solver-augmented dataset", "Existing tests cover many import/action guardrails but not a broad solver-grounded golden set.", "Log validated traces for regression first; fine-tune only after measured need."),
    ]
    for values in rows:
        cells = table.add_row().cells
        for i, text in enumerate(values):
            p = cells[i].paragraphs[0]; set_font(p.add_run(text), 9.2, INK)
    set_table_geometry(table, [2500, 3300, 3560])

    heading(doc, "Recommended target architecture", 1)
    add_numbered(doc, [
        "Canonicalize the hand into PokerStateV1: variant, seats, positions, stacks, blinds/antes, cards, board, full action history, decision node, and derived legal actions.",
        "Run SolverStatePreflight: deck uniqueness, street/action order, pot/stack conservation, amount-to-call, legal actions, normalized ranges, and complete solver-node traversal.",
        "Build PokerEvidencePacketV1: provenance/confidence, pot odds, equity method/sample count, range weights/version, solver strategy and configuration, accuracy/status, and limitations.",
        "Apply DecisionContractV1: preserve the full mixed strategy; use a declared selection rule only when the UI needs one action; explicitly label solver versus fallback modes.",
        "Generate a structured explanation with evidence references, then run semantic validation for legality, numeric grounding, action/mix agreement, and unsupported precision.",
    ])
    add_callout(doc, "Core invariant", "For one analysis ID, the UI, solver, recommendation, explanation, and audit log must all reference the same state hash and evidence-packet version.", "F4F6F9")

    heading(doc, "Evidence packet: minimum useful shape", 1)
    add_bullets(doc, [
        "state — canonical cards, positions, stacks, betting history, decision node, and legal actions.",
        "provenance — manual/imported/derived source, confidence, repairs, parser version, and state hash.",
        "range — weighted player ranges, blockers removed, source, abstraction, and version.",
        "math — pot, call price, pot odds, equity, sampling/algorithm details, and EV definitions.",
        "strategy — action sizes/frequencies, node path, solver/version/configuration, convergence/accuracy, and action EVs when supported.",
        "limitations — ambiguity, unsupported node, timeout, multiway approximation, missing stack, or range sensitivity.",
    ])

    heading(doc, "Highest-priority controls", 1)
    heading(doc, "1. Solver state preflight", 2)
    add_p(doc, "Fail closed when the solver node cannot be reproduced exactly. The current adapter can traverse a partial action path and still summarize the last matched node; the safer contract is to require complete traversal. Derive legal actions and amount-to-call on the server rather than trusting presentation state.")
    heading(doc, "2. Deterministic recommendation arbitration", 2)
    add_p(doc, "When a supported, validated solver result exists, make its mixed strategy the primary decision. The LLM should receive the action/mix as immutable evidence. When the solver is unavailable, label the result Calculated fallback or Incomplete state and never describe it as GTO.")
    heading(doc, "3. Semantic explanation validation", 2)
    add_p(doc, "Extend beyond truncation, repetition, and required headings. Parse structured output and verify legal actions, recommendation/mix agreement, every numeric claim, cited range buckets, and the distinction between exact and qualitative blocker claims.")

    heading(doc, "Methods worth adapting—without training first", 1)
    heading(doc, "Regret as a coaching signal", 2)
    add_p(doc, "The paper uses CFR-derived regret as dense training feedback. PokerCoach can turn the same idea into a user feature: compare the recorded action with the best action and show EV loss/regret in big blinds, with an uncertainty band. This lets users prioritize costly leaks instead of treating every frequency deviation equally. Do not infer regret from action frequency; require action EVs or a solver-supported proxy.")
    heading(doc, "HR / FA / AC as an evaluation rubric", 2)
    add_bullets(doc, [
        "Factual alignment (FA): cards, pot, call price, equity, ranges, and solver values match packet fields.",
        "Action–reasoning consistency (AC): recommendation is legal and matches the decision contract; candidates match the legal set.",
        "Heuristic reasoning (HR): explanation uses decision-relevant quantities/ranges, avoids hand-strength slogans, and states uncertainty.",
    ])
    add_p(doc, "Make FA and AC deterministic. Use an LLM judge only for the residual qualitative portion of HR, and treat it as diagnostic rather than ground truth.")
    heading(doc, "Solver-augmented examples", 2)
    add_p(doc, "Begin by logging state → evidence → decision → explanation → validation records. Use them as regression fixtures, prompt examples, and model-comparison data. Fine-tuning becomes rational only if telemetry shows persistent failures that constrained generation cannot solve economically.")

    heading(doc, "Roadmap and acceptance gates", 1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for i, text in enumerate(["Phase", "Deliverables", "Exit criteria"]):
        shade(table.rows[0].cells[i], PALE_BLUE)
        p = table.rows[0].cells[i].paragraphs[0]; set_font(p.add_run(text), 9.5, INK, True)
    roadmap = [
        ("1 • 1–2 weeks", "Schemas; unified server computation; state hash/provenance; semantic gates; mode badge.", "Zero illegal recommendations; zero valid-solver/action disagreements; every number traceable."),
        ("2 • 2–4 weeks", "Solver in main orchestration; full mix; action EV/regret where available; strict explanation schema; cache by state/range/solver version.", "≥99% validation after one retry on supported spots; correct fallback labels; latency/failures observable."),
        ("3 • 2–3 weeks", "Stratified golden set; PokerBench feasibility check; current-vs-new ablations; validated trace curation.", "Decision accuracy, EV loss, FA, AC, illegal rate, unsupported precision, latency, and cache hit tracked."),
        ("4 • later", "SFT/RL decision only after sufficient clean traces and a measured unresolved issue.", "Training business case beats prompting/constrained generation on quality, latency, or cost."),
    ]
    for vals in roadmap:
        cells = table.add_row().cells
        for i, text in enumerate(vals):
            p = cells[i].paragraphs[0]; set_font(p.add_run(text), 9.2, INK)
    set_table_geometry(table, [1500, 4200, 3660])

    heading(doc, "Product experiments", 1)
    add_numbered(doc, [
        "Why this mix? Compare solver mix + evidence-cited explanation with today’s independent commentary; measure factual errors, trust, and time-to-understanding.",
        "Leak severity. Show EV loss/regret for the recorded action; measure review of high-cost errors and performance on similar spots.",
        "Uncertainty-aware import. Ask users to correct only fields blocking solver preflight; measure correction completion and solver-success lift.",
        "Explanation depth. Offer concise, study, and advanced views from the same packet; the decision remains unchanged.",
        "Off-equilibrium sensitivity. Perturb opponent range weights and sizes; show whether the recommendation is robust or assumption-sensitive.",
    ], style="List Number 2")

    heading(doc, "Critical limitations", 1)
    add_bullets(doc, [
        "The paper mainly studies Leduc and heads-up limit hold’em, not six-max NLHE. Transfer the architecture, not the headline chip totals.",
        "Matched results use 100 games. Poker variance makes chip totals directional rather than precise win-rate estimates.",
        "GTO is conditional on game model, ranges, action tree, convergence, and correct state. PokerCoach’s adapter currently fixes bet/raise sizes, approximates effective stack as four times pot, and uses low default iterations; do not market this as exact GTO.",
        "Multiway NLHE and off-tree actions require explicit approximation policies. The paper’s simpler three-player result does not validate multiway NLHE solving.",
        "Reasoning scores partly rely on an LLM judge; replace factual and alignment judgments with deterministic tests wherever possible.",
        "Position the product for post-hand study rather than real-time assistance, with responsible-gambling safeguards.",
    ])
    add_callout(doc, "Decision", "Proceed with unified evidence and semantic decision contracts. Defer fine-tuning. Treat solver fidelity, state validation, and evaluation infrastructure as prerequisites.")

    heading(doc, "Sources", 1)
    sources = [
        ("Lin et al. (2026), ToolPoker paper", "https://arxiv.org/html/2602.00528"),
        ("Zinkevich et al. (2007), Counterfactual Regret Minimization", "https://papers.nips.cc/paper_files/paper/2007/file/08d98638c6fcd194a4b1e6992063e944-Paper.pdf"),
        ("Brown et al. (2019), Deep Counterfactual Regret Minimization", "https://proceedings.mlr.press/v97/brown19b.html"),
        ("Zhuang et al. (2025), PokerBench paper", "https://ojs.aaai.org/index.php/AAAI/article/download/34814/36969"),
        ("PokerBench official repository and dataset description", "https://github.com/pokerllm/pokerbench"),
    ]
    for label, url in sources:
        p = doc.add_paragraph(style="List Bullet"); add_link(p, label, url)
    add_p(doc, "Repository evidence inspected 25 August 2026: public/app.js; src/analysis/pipeline.js; src/solver/texasSolverAdapter.js; src/analysis/strategicOutputValidation.js.")

    heading(doc, "Research scope and confidence", 1)
    add_p(doc, "High confidence: architectural recommendation, state/action alignment priorities, and current-code mapping. Medium confidence: exact implementation effort and action-EV availability in the present TexasSolver integration. Low confidence: direct transfer of the paper’s gameplay gains to six-max NLHE. Research stopped after the paper’s full method/ablations/errors were checked, the primary CFR/DeepCFR sources and PokerBench were reviewed, and further search was unlikely to change the product recommendation.")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
